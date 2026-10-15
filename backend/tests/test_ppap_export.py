"""AC08/AC10/AC16: inspect actual exported workbook cells and saved layout.

All values below are explicit test fixtures, never claimed as factory evidence.
"""
from __future__ import annotations

import re
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from app.services import ppap_export, ppap_storage, ppap_template


def export_fixture() -> dict:
    """Small, deliberately identified fixture spanning engineering and raw data."""
    def element(element_id: str, **content) -> dict:
        return {"element_id": element_id, "requirement": "submit", "status": "DRAFTED",
                "status_reason": "fixture; pending human review", "applicability_basis": "test fixture only",
                "content": {"template_source": "project_defined", "approval_state": "DRAFT", **content},
                "source_refs": {"input_ids": ["INPUT-FIXTURE-L-001"]}, "check_result": {"findings": []}}
    operation = {"part_side": "LEFT", "op_id": "OP-FIXTURE-10", "op_name": "fixture confirmed operation", "station": "fixture station"}
    char = {"part_side": "LEFT", "characteristic_id": "CHAR-FIXTURE-L-01", "name": "fixture dimension",
            "special": True, "unit": "mm", "usl": 10.5, "lsl": 9.5, "nominal": 10.0}
    task = {"task_id": "PPAP-TEST-EXPORT", "customer_name": "TEST FIXTURE - NOT REAL FACTORY",
            "part_no": "FIXTURE-L", "part_name": "test fixture left part", "part_side": "LEFT",
            "drawing_version": "TEST-R1", "template_source": "project_defined",
            "package_state": "NOT_SUBMITTED", "customer_state": "NOT_SUBMITTED"}
    elements = [element("process_flow_diagram", operations=[operation], relations=[{"from_op": "OP-FIXTURE-10", "to_op": "OP-FIXTURE-20"}]),
                element("pfmea", operations=[operation], characteristics=[char], rows=[{"op_id": "OP-FIXTURE-10", "failure_mode": None, "severity": None}],
                        pending_confirmations=["Risk ratings need engineering confirmation"]),
                element("control_plan", operations=[operation], characteristics=[char], rows=[{"characteristic_id": char["characteristic_id"], "measurement_method": None}]),
                element("dimensional_results", rows=[{"part_side": "LEFT", "point": "POINT-FIXTURE-1", "value": 10.1, "usl": 10.5, "lsl": 9.5,
                                                           "unit": "mm", "result": "OK", "source_ref": "test fixture measurement record"},
                                                          {"point": "POINT-FIXTURE-2", "value": None, "result": "MISSING", "result_basis": "no original value"}]),
                element("msa", studies=[{"study_id": "MSA-FIXTURE-L", "part_side": "LEFT", "characteristic_id": char["characteristic_id"],
                                         "raw_data": {"unit": "mm", "measurements": [
                                             {"appraiser": "FIXTURE-A", "parts": [{"part": "P1", "trials": [1.0, 1.1]}, {"part": "P2", "trials": [2.0, 2.1]}]},
                                             {"appraiser": "FIXTURE-B", "parts": [{"part": "P1", "trials": [1.1, 1.2]}, {"part": "P2", "trials": [2.1, 2.2]}]}]},
                                         "result": {"state": "OK", "value": {"grr": 0.1, "percent_grr": 20.0}}}],
                        calculations=[{"metric": "percent_grr", "value": 20.0, "acceptance_state": "NEEDS_CONFIRMATION", "basis": "no customer acceptance rule"}]),
                element("initial_process_study", studies=[{"study_id": "PROCESS-FIXTURE-L", "part_side": "LEFT", "raw_data": {"values": [9.8, 10.0, 10.2], "unit": "mm"},
                                                           "result": {"value": {"mean": 10.0, "cp": 0.94}, "state": "INCOMPLETE"}}])]
    evidence = [{"evidence_id": "EVIDENCE-FIXTURE-1", "source_type": "user_confirmation", "source_id": "SOURCE-FIXTURE-1",
                 "source_version": "TEST-R1", "locator": "test fixture only", "content_hash": "0" * 64,
                 "read_status": "OK", "review_state": "CONFIRMED", "element_ids": ["process_flow_diagram"], "summary": "test fixture evidence"},
                {"evidence_id": "EVIDENCE-FIXTURE-ERROR", "source_type": "openmes_inspection_record", "source_id": "SOURCE-FIXTURE-ERROR",
                 "read_status": "ERROR", "error": {"message": "fixture transport error, not an empty result"}}]
    issues = [{"issue_id": "ISSUE-FIXTURE-1", "check_id": "PPAP-CHK-MSA-001", "element_id": "msa", "kind": "RULE_UNCONFIRMED",
               "severity": "BLOCKER", "status": "OPEN", "blocks_approval": True, "title": "MSA acceptance rule unconfirmed",
               "detail": "No acceptance rules supplied; cannot claim approved.", "basis": "test fixture", "responsible_role": "quality_manager"}]
    return {"task": task, "elements": elements, "evidence": evidence, "issues": issues,
            "version": 1, "snapshot_hash": "1" * 64, "compile_run_id": "RUN-TEST-EXPORT-1"}


@pytest.fixture
def output(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    package = ppap_export.export_package(**export_fixture())
    path = ppap_storage.storage_root() / package["rel_path"]
    return package, path


def sheet_values(ws):
    return [cell.value for row in ws for cell in row if cell.value is not None]


def test_export_actual_content_and_provenance(output):
    package, path = output
    wb = load_workbook(path, data_only=False)
    assert len(wb.sheetnames) == len(ppap_template.ELEMENTS) + 4
    assert all(s.sheet_name in wb.sheetnames for s in ppap_template.ELEMENTS)
    assert "TEST FIXTURE - NOT REAL FACTORY" in sheet_values(wb["资料包概览"])
    assert "FIXTURE-L" in sheet_values(wb["资料包概览"])
    assert "OP-FIXTURE-10" in sheet_values(wb["过程流程图"])
    assert "CHAR-FIXTURE-L-01" in sheet_values(wb["过程FMEA"])
    assert "CHAR-FIXTURE-L-01" in sheet_values(wb["控制计划"])
    assert "EVIDENCE-FIXTURE-1" in sheet_values(wb["来源索引"])
    assert "ERROR" in sheet_values(wb["来源索引"])
    assert "fixture transport error, not an empty result" in sheet_values(wb["来源索引"])
    assert "MSA acceptance rule unconfirmed" in sheet_values(wb["缺项与冲突"])
    assert package["metadata"]["template_source"] == "project_defined"
    assert package["metadata"]["is_customer_template"] is False
    assert package["byte_size"] == path.stat().st_size
    assert ppap_storage.sha256_bytes(path.read_bytes()) == package["content_hash"]
    hash_cells = [c for row in wb["资料包概览"] for c in row if c.value == "1" * 64]
    assert len(hash_cells) == 1 and hash_cells[0].data_type == "s" and hash_cells[0].number_format == "@"
    wb.close()


def test_raw_values_not_replaced_by_approved_or_missing_zero(output):
    _, path = output
    wb = load_workbook(path, data_only=False)
    sizes = sheet_values(wb["尺寸结果"])
    assert 10.1 in sizes and "MISSING" in sizes
    msa = sheet_values(wb["MSA"])
    assert "FIXTURE-A" in msa and "FIXTURE-B" in msa
    assert all(v in msa for v in (1.0, 1.1, 1.2, 2.0, 2.1, 2.2))
    assert "NEEDS_CONFIRMATION" in msa
    assert "Risk ratings need engineering confirmation" in sheet_values(wb["过程FMEA"])
    assert "APPROVED" not in sheet_values(wb["资料包概览"])
    # Blank PFMEA ratings stay blank and are not borrowed from a historical sample.
    ws = wb["过程FMEA"]
    for row in ws:
        for cell in row:
            if cell.value == "严重程度":
                assert ws.cell(cell.row + 1, cell.column).value is None
    wb.close()


def test_saved_formulas_have_exact_raw_sample_range(output):
    package, path = output
    wb = load_workbook(path, data_only=False)
    ws = wb["初始过程研究"]
    formulas = [c.value for row in ws for c in row if c.data_type == "f"]
    assert len(formulas) == 4
    assert any(f.startswith("=AVERAGE(") for f in formulas)
    assert any(f.startswith("=STDEV(") for f in formulas)
    for formula in formulas:
        start, end = map(int, re.search(r"B(\d+):B(\d+)", formula).groups())
        assert [ws.cell(row, 2).value for row in range(start, end + 1)] == [9.8, 10.0, 10.2]
    assert package["metadata"]["workbook_inspection"]["formula_count"] == 4
    assert package["metadata"]["workbook_inspection"]["formula_cache_recalculated"] is False
    wb.close()


def test_saved_print_layout_editable_and_all_sheets_marked(output, tmp_path):
    _, path = output
    wb = load_workbook(path, data_only=False)
    for ws in wb:
        assert ws["A2"].value == ppap_template.SHEET_BANNER
        assert ws.page_setup.orientation == "landscape"
        assert int(ws.page_setup.paperSize) == int(ws.PAPERSIZE_A3)
        assert ws.page_setup.fitToWidth == 1
        assert ws.page_setup.fitToHeight == 0
        assert ws.print_title_rows == "$1:$4"
        assert ws.freeze_panes == "A5"
        assert f"$H${ws.max_row}" in str(ws.print_area)
        assert not ws.protection.sheet
        assert all(d.width >= 17 for d in ws.column_dimensions.values())
    wb["资料包概览"]["C5"] = "OFFLINE EDIT TEST"
    edited = tmp_path / "edited-fixture.xlsx"
    wb.save(edited)
    reopened = load_workbook(edited)
    assert reopened["资料包概览"]["C5"].value == "OFFLINE EDIT TEST"
    reopened.close()
    wb.close()


@pytest.mark.parametrize("value", ['=HYPERLINK("https://untrusted.invalid","click")', "+cmd|' /C calc'!A0", "-1+1", "@SUM(A1:A2)", "  =1+1"])
def test_nested_text_formula_injection_is_literal(value, tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    fixture["elements"][0]["content"]["operations"][0]["op_name"] = value
    result = ppap_export.export_package(**fixture)
    wb = load_workbook(ppap_storage.storage_root() / result["rel_path"])
    cells = [c for row in wb["过程流程图"] for c in row if c.value == value]
    assert len(cells) == 1
    assert cells[0].data_type == "s" and cells[0].quotePrefix
    assert cells[0].hyperlink is None
    wb.close()


def test_existing_artifact_cannot_be_overwritten(output):
    package, path = output
    original = path.read_bytes()
    with pytest.raises(FileExistsError):
        ppap_export.export_package(**export_fixture())
    assert path.read_bytes() == original
    assert ppap_storage.sha256_bytes(original) == package["content_hash"]


@pytest.mark.parametrize("field,value", [("task_id", "../escape"), ("task_id", "C:/escape"), ("compile_run_id", "../../outside")])
def test_artifact_path_escape_rejected(field, value, tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    if field == "task_id":
        fixture["task"][field] = value
    else:
        fixture[field] = value
    with pytest.raises(ValueError):
        ppap_export.export_package(**fixture)
    assert not list(tmp_path.rglob("*.xlsx"))


def test_unmapped_customer_template_not_claimed(tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    fixture["task"]["template_source"] = "customer_template"
    with pytest.raises(ValueError, match="正式客户模板"):
        ppap_export.export_package(**fixture)


def test_credentials_refused_without_echo(tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    fixture["evidence"][0]["error"] = {"authorization": "FAKE CREDENTIAL SENTINEL"}
    with pytest.raises(ValueError) as error:
        ppap_export.export_package(**fixture)
    assert "FAKE CREDENTIAL SENTINEL" not in str(error.value)
    assert not list(tmp_path.rglob("*.xlsx"))


def test_long_source_text_preserved_without_clipped_single_row(tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    long = "采购证明完整内容" * 200
    fixture["evidence"][0]["summary"] = long
    result = ppap_export.export_package(**fixture)
    wb = load_workbook(ppap_storage.storage_root() / result["rel_path"])
    parts = [v for v in sheet_values(wb["来源索引"]) if isinstance(v, str) and "采购证明" in v]
    assert "".join(parts) == long
    assert max(d.height or 0 for d in wb["来源索引"].row_dimensions.values()) <= 360
    wb.close()


def test_export_does_not_embed_external_links_macros_or_source_images(output):
    _, path = output
    with ZipFile(path) as archive:
        assert not any(p.startswith("xl/externalLinks/") for p in archive.namelist())
        assert not any("vbaProject" in p or p.startswith("xl/embeddings/") or p.startswith("xl/media/") for p in archive.namelist())


def test_same_planned_owner_reviewer_warns_without_signing_or_approving(tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    fixture["task"].update(owner="fixture-actor", reviewer="fixture-actor")
    result = ppap_export.export_package(**fixture)
    wb = load_workbook(ppap_storage.storage_root() / result["rel_path"])
    values = sheet_values(wb["资料包概览"])
    assert any(isinstance(v, str) and "任务负责人和指定审核人为同一账号" in v for v in values)
    assert "DRAFT" in values and "APPROVED" not in values
    wb.close()


@pytest.mark.parametrize("version", [True, 0, -1, "1"])
def test_invalid_version_rejected(version, tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    fixture["version"] = version
    with pytest.raises(ValueError):
        ppap_export.export_package(**fixture)


def test_unknown_or_duplicate_element_rejected(tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path / "ppap_store"))
    fixture = export_fixture()
    fixture["elements"].append({"element_id": "invented_unregistered_element"})
    with pytest.raises(ValueError):
        ppap_export.export_package(**fixture)
    fixture = export_fixture()
    fixture["elements"].append(fixture["elements"][0])
    with pytest.raises(ValueError):
        ppap_export.export_package(**fixture)
