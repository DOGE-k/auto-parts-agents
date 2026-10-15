"""Regression checks for mixed-language layout and honest formula provenance."""
from io import BytesIO
import math
import statistics

from openpyxl import Workbook, load_workbook
import pytest

from app.services import ppap_export, ppap_storage
from test_ppap_export import export_fixture


@pytest.mark.parametrize("text", ["中文说明" * 300, "LONG_IDENTIFIER_0123456789/" * 200,
                                  "第一行\n第二行\r\n" * 100, "\n" * 200, "\t数据ＭＷ@" * 300])
def test_layout_chunks_preserve_all_characters_and_excel_row_limit(text):
    chunks = ppap_export._text_chunks(text, 1)
    assert "".join(chunks) == text
    assert len(chunks) > 1
    assert all(ppap_export._height(chunk, 1) <= 360 for chunk in chunks)


def test_cjk_and_explicit_line_breaks_increase_required_height():
    assert ppap_export._height("中" * 52, 1) >= 134
    assert ppap_export._height("中" * 52, 1) > ppap_export._height("x" * 52, 1)
    assert ppap_export._height("中" * 102, 1) >= 242
    assert ppap_export._height("a\nb\nc", 1) >= 62


def test_saved_table_height_preserves_long_text_and_existing_column_width(tmp_path):
    wb = Workbook()
    ws = wb.active
    writer = ppap_export._SheetWriter(ws, "布局回归测试", {"task_id": "TEST", "data_mode": "TEST"}, 1)
    chinese = "中" * 52
    mixed = "外观批准报告（AAR）的适用性未确认。appearance_approval_report 的提交要求尚未确认，建议值 submit，外观件且客户要求时提交。" + "补充说明" * 30
    writer.table([{"a": "one", "b": chinese}, {"a": "two", "b": mixed}])
    writer.finish()
    path = tmp_path / "layout.xlsx"
    wb.save(path)
    saved = load_workbook(path)
    sheet = saved.active
    chinese_cells = [c for row in sheet for c in row if c.value == chinese]
    assert len(chinese_cells) == 1
    assert sheet.row_dimensions[chinese_cells[0].row].height >= 134
    start = next(c.row for row in sheet for c in row if c.value == "two")
    assert "".join(sheet.cell(r, 2).value or "" for r in range(start, sheet.max_row + 1)) == mixed
    assert max(d.height for d in sheet.row_dimensions.values()) <= 360
    assert all(d.width == 17 for d in sheet.column_dimensions.values())
    saved.close()


def test_notes_fields_titles_and_table_headers_use_same_bounded_layout():
    wb = Workbook()
    ws = wb.active
    writer = ppap_export._SheetWriter(ws, "布局回归测试", {}, 1)
    jobs = [("note", "\n".join(["备注测试"] * 100)),
            ("field", "\n".join(["字段测试"] * 100)),
            ("section", "\n".join(["标题测试"] * 100)),
            ("header", "\n".join(["列标题测试"] * 100))]
    for kind, text in jobs:
        start = writer.row
        if kind == "note":
            writer._note(text)
        elif kind == "field":
            writer.fields({"字段标签": text})
        elif kind == "section":
            writer._section(text)
        else:
            writer._table_header([text])
        column = 3 if kind == "field" else 1
        assert "".join(ws.cell(r, column).value or "" for r in range(start, writer.row)) == text
    assert max(d.height for d in ws.row_dimensions.values()) <= 360
    wb.close()


@pytest.mark.parametrize("mode", ["REAL", "TEST"])
def test_data_mode_exported_visibly_and_in_metadata(mode, tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path))
    fixture = export_fixture()
    fixture["task"]["data_mode"] = mode
    result = ppap_export.export_package(**fixture)
    wb = load_workbook(ppap_storage.storage_root() / result["rel_path"])
    assert result["metadata"]["data_mode"] == mode
    assert result["metadata"]["test_package"] is (mode == "TEST")
    for ws in wb:
        assert ("测试资料包" in ws["A1"].value) is (mode == "TEST")
        if mode == "TEST":
            assert ws["A4"].value == ppap_export._TEST_BANNER
            assert ws.oddFooter.right.text == "测试资料包（TEST）"
        assert not ws.protection.sheet
    assert ("测试资料包" in wb.properties.title) is (mode == "TEST")
    assert not any(c.value == "APPROVED" for row in wb["资料包概览"] for c in row)
    wb.close()


def test_auxiliary_formulas_have_independent_expected_values_without_claiming_native_recalc(tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path))
    result = ppap_export.export_package(**export_fixture())
    inspection = result["metadata"]["workbook_inspection"]
    assert inspection["auxiliary_formula_source_verified"] is True
    assert inspection["formula_cache_recalculated"] is False
    assert inspection["native_engine_verified"] is False
    expected = [statistics.fmean([9.8, 10.0, 10.2]), 9.8, 10.2, statistics.stdev([9.8, 10.0, 10.2])]
    assert [c["independent_expected_value"] for c in inspection["auxiliary_formula_checks"]] == pytest.approx(expected)
    assert all(c["sample_count"] == 3 for c in inspection["auxiliary_formula_checks"])
    path = ppap_storage.storage_root() / result["rel_path"]
    cached = load_workbook(path, data_only=True)
    for item in inspection["auxiliary_formula_checks"]:
        assert cached[item["sheet"]][item["cell"]].value is None
    cached.close()
    wb = load_workbook(path)
    item = inspection["auxiliary_formula_checks"][0]
    start = int(item["source_range"].split(":")[0][1:])
    wb[item["sheet"]].cell(start, 2).value = "missing measurement"
    damaged = BytesIO()
    wb.save(damaged)
    wb.close()
    with pytest.raises(ValueError, match="完整有限数值样本"):
        ppap_export.validate_workbook(damaged.getvalue())


def test_unknown_data_mode_refused_before_file_creation(tmp_path, monkeypatch):
    monkeypatch.setenv("PPAP_STORAGE_DIR", str(tmp_path))
    fixture = export_fixture()
    fixture["task"]["data_mode"] = "PRODUCTION_ASSUMED"
    with pytest.raises(ValueError, match="数据模式"):
        ppap_export.export_package(**fixture)
    assert not list(tmp_path.rglob("*.xlsx"))
