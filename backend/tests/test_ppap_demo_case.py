"""AC08/AC09/AC10/AC12：端到端生成一份可编辑 PPAP 包的回归用例。

本用例把 scripts/ppap_demo_case.py 的案例固化为断言，确保"一口气生成文档"
这条路不会被后续改动打断。所有数据都是显式测试标记（DEMO- 前缀、TEST 模式），
不关联真实 ERP/MES，也不冒充真实客户提交。
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND.parent / "scripts"))

from app.persistence.models import AuditRecordRow  # noqa: E402
from app.persistence.database import SessionLocal  # noqa: E402
from app.services import ppap, ppap_template  # noqa: E402


@pytest.fixture(scope="module")
def demo():
    """跑一次完整流程，供本模块所有断言复用（生成一次产物，避免重复编译）。"""
    import ppap_demo_case

    result = ppap_demo_case.build()
    assert result["status"] == "COMPILED", result
    return result


def test_case_reaches_compiled_without_blockers(demo):
    task = ppap.get_task(demo["task_id"], include_content=False)
    # 归档后包审核状态为 ARCHIVED；运行状态在归档时不再宣称 COMPILED。
    assert task["task"]["package_state"] == "ARCHIVED"
    assert task["task"]["data_mode"] == "TEST"
    versions = ppap.list_versions(demo["task_id"])["items"]
    assert len(versions) == 1 and versions[0]["status"] == "ARCHIVED"
    assert versions[0]["frozen_at"]


def test_all_22_elements_are_drafted_with_base_info(demo):
    detail = ppap.get_task(demo["task_id"], include_content=True)
    assert len(detail["elements"]) == 22
    for element in detail["elements"]:
        content = element["content"] or {}
        assert content["base_info"]["part_no"] == "DEMO-BD-2401-L"
        assert content["template_source"] == "project_defined"
        assert content["approval_state"] == "DRAFT"


def test_special_characteristic_propagates_to_control_plan_and_pfmea(demo):
    detail = ppap.get_task(demo["task_id"], include_content=True)
    by_id = {item["element_id"]: item for item in detail["elements"]}
    for element_id in ("control_plan", "pfmea"):
        ids = {row["characteristic_id"] for row in by_id[element_id]["content"]["characteristics"]}
        assert "DEMO-CHAR-SPC-01" in ids
    for element_id in ("process_flow_diagram", "pfmea", "control_plan"):
        ops = {row["op_id"] for row in by_id[element_id]["content"]["operations"]}
        assert ops == {op["op_id"] for op in ppap_demo_case_ops()}


def ppap_demo_case_ops():
    import ppap_demo_case

    return ppap_demo_case.OPERATIONS


def test_excel_has_real_content_and_editable_print_layout(demo):
    reading = ppap.read_artifact(demo["task_id"], demo["artifact_id"])
    assert reading["ok"]
    wb = load_workbook(io.BytesIO(reading["content"]))
    assert len(wb.sheetnames) == 26
    for required in ("资料包概览", "控制计划", "尺寸结果", "MSA", "缺项与冲突", "来源索引"):
        assert required in wb.sheetnames
    # 测试标识必须写在表内，避免被当成真实资料包
    banner = " ".join(
        str(cell.value) for row in wb["控制计划"].iter_rows(min_row=1, max_row=4)
        for cell in row if cell.value is not None
    )
    assert "项目自定义模板" in banner
    assert "TEST" in banner or "测试" in banner
    # 公式必须引用真实原始样本区间，而不是写死结果
    formulas = [cell.value for ws in wb.worksheets for row in ws.iter_rows()
                for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert any(value.startswith("=STDEV(") for value in formulas)
    assert any(value.startswith("=AVERAGE(") for value in formulas)
    for ws in wb.worksheets:
        assert ws.page_setup.orientation == "landscape"
        assert ws.print_title_rows == "$1:$4"
        assert ws.print_area


def test_default_deployment_has_no_confirmed_rules(tmp_path, monkeypatch):
    """未配置 acceptance_rules.json 时，指标必须是「规则未确认」而不是「合格」。"""
    from app.services import ppap_calc

    monkeypatch.setenv("PPAP_RULES_FILE", str(tmp_path / "absent.json"))
    rules = ppap_calc.load_acceptance_rules()
    assert rules["configured"] is False
    outcome = ppap_calc.classify_metric("percent_grr", 12.0, rules)
    assert outcome["acceptance_state"] == "NEEDS_CONFIRMATION"
    assert outcome["reference_only"] is True


def test_demo_does_not_write_audit_table(demo):
    with SessionLocal() as session:
        assert session.query(AuditRecordRow).count() == 0


def test_second_compile_is_idempotent(demo):
    replay = ppap.compile_task(demo["task_id"], actor="demo-author")
    assert replay["success"] and replay["idempotent_replay"]
    assert replay["artifact"]["artifact_id"] == demo["artifact_id"]
