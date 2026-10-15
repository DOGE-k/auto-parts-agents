"""Editable PPAP snapshot export using the project's own, non-customer template.

This module reads already-authorized compilation facts. It does not query or write
ERP/MES, update review state, fabricate measurements or sign a document. Sources
and unresolved findings remain visible in the exported workbook. Native customer
templates/images are not available in this implementation and are not claimed.
"""
from __future__ import annotations

from collections.abc import Mapping, Sequence
from io import BytesIO
import math
import re
import statistics
import unicodedata
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

from app.services import ppap_storage, ppap_template

MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXPORT_VERSION = "ppap-internal-xlsx-1.1"
MAX_COLUMNS = 8
_COLUMN_WIDTH = 17
# Leave space below Excel's 409.5 pt limit and for font/rendering differences.
_MAX_ROW_HEIGHT = 360
_TEST_BANNER = "测试资料包（TEST）：仅供功能验收，不能作为真实生产证据或正式客户提交资料。"
_TEAL = "076C68"
_NAVY = "16323D"
_PALE = "E7F2F1"
_ORANGE = "FFF1D7"
_GRAY = "667A84"
_LINE = "D5E3E7"
_FONT = "Microsoft YaHei"
_SECRET_KEYS = {
    "token", "access_token", "refresh_token", "bearer", "bearer_token",
    "password", "api_key", "api_secret", "secret", "cookie", "cookies",
    "authorization", "real_write_token", "project_auth_password",
}
_LABELS = {
    "task_id": "任务编号", "data_mode": "数据模式", "customer_id": "客户编号", "customer_name": "客户",
    "part_no": "零件号", "part_name": "零件名称", "part_side": "左右属性",
    "drawing_version": "图纸版本", "engineering_change_level": "工程变更等级",
    "submission_reason": "提交原因", "submission_level": "提交等级",
    "template_id": "模板编号", "template_version": "模板版本",
    "template_source": "模板来源", "owner": "负责人", "reviewer": "指定审核人",
    "due_date": "期望完成日期", "run_state": "运行状态", "package_state": "包审核状态",
    "customer_state": "客户处理状态", "business_object": "业务关联",
    "base_info": "统一基础信息", "operations": "工序清单", "characteristics": "特性清单",
    "relations": "工序关系", "rows": "记录明细", "studies": "原始研究与计算",
    "calculations": "确定性指标与规则状态", "references": "关联证据",
    "pending_confirmations": "待确认事项", "approval_state": "编制状态",
    "rule_hash": "规则指纹", "rule_version": "规则版本", "op_id": "工序编号",
    "op_no": "工序号", "op_name": "工序名称", "station": "工位", "sequence": "顺序",
    "description": "描述", "from_op": "前工序", "to_op": "后工序",
    "relation_type": "关系类型", "characteristic_id": "特性编号", "name": "名称",
    "characteristic_type": "特性类型", "special": "特殊特性", "unit": "单位",
    "usl": "规格上限", "lsl": "规格下限", "nominal": "名义值", "target": "目标值",
    "drawing_location": "图纸位置", "point": "测点编号", "point_id": "测点编号",
    "value": "数值", "sample_no": "样本号", "measured_by": "原始记录测量人",
    "measured_at": "原始记录测量时间", "missing_reason": "缺失说明",
    "source_ref": "来源引用", "source_refs": "来源引用", "source_type": "来源类型",
    "source_id": "来源记录编号", "source_version": "来源版本", "locator": "定位信息",
    "content_hash": "内容指纹", "read_status": "读取状态", "review_state": "证据确认状态",
    "evidence_id": "证据编号", "element_id": "条目编号", "element_ids": "关联条目",
    "study_id": "研究编号", "raw_data": "原始数据", "measurements": "MSA原始测量",
    "values": "过程研究原始样本", "subgroups": "原始子组", "method": "方法",
    "tolerance": "公差", "result": "计算/比对结果", "result_basis": "依据",
    "metric": "指标", "calc_state": "计算状态", "state": "状态",
    "acceptance_state": "规则判定状态", "threshold": "已确认阈值",
    "rule_source": "规则来源", "basis": "依据", "flags": "边界标记",
    "notes": "说明", "note": "备注", "problems": "计算问题", "summary": "摘要",
    "ev": "重复性 EV", "av": "再现性 AV", "grr": "GRR", "pv": "零件间变差 PV",
    "tv": "总变差 TV", "percent_ev": "重复性占比（%）", "percent_av": "再现性占比（%）",
    "percent_grr": "GRR占比（%）", "percent_tolerance": "公差占比（%）",
    "ndc": "可分辨类别数", "ndc_truncated": "可分辨类别数（取整）",
    "cp": "Cp", "cpk": "Cpk", "pp": "Pp", "ppk": "Ppk", "count": "样本数量",
    "mean": "平均值", "min": "最小值", "max": "最大值",
    "stddev_within": "组内标准差", "stddev_overall": "总体样本标准差",
    "failure_mode": "潜在失效模式", "effect": "失效影响", "cause": "潜在原因",
    "prevention_control": "预防措施", "detection_control": "探测措施",
    "severity": "严重程度", "occurrence": "发生度", "detection": "探测度",
    "rpn": "RPN", "risk_method": "风险评级方法", "measurement_method": "测量方法",
    "sample_size": "样本容量", "frequency": "频次", "reaction_plan": "反应计划",
    "supplier": "供应商", "lot_no": "材料批次", "report_no": "报告编号",
    "test_item": "试验项目", "spec_ref": "规格依据", "item": "项目", "quantity": "数量",
    "change_description": "变更描述", "change_no": "变更编号", "change_date": "变更日期",
    "level": "等级", "reason": "原因", "requirement_text": "客户要求", "clause": "条款",
    "record_type": "实物记录类型", "storage_location": "实物存放位置",
    "photo_refs": "照片引用", "instruction": "要求说明", "stage": "阶段",
    "findings": "检查发现", "input_ids": "结构化输入编号", "evidence_ids": "证据编号",
    "issue_id": "待办编号", "check_id": "检查编号", "kind": "问题类型", "title": "标题",
    "detail": "详细说明", "required_evidence": "所需证据", "responsible_role": "责任岗位",
    "status": "处理状态", "blocks_approval": "是否阻断批准", "resolution_note": "处理说明",
    "resolved_by": "处理账号", "resolved_at": "处理时间", "created_at": "创建时间",
    "updated_at": "更新时间", "truncated": "读取不完整", "error": "读取错误",
    "message": "错误说明", "dedup_key": "唯一发现标识", "field_path": "受影响字段",
}
_REQUIREMENTS = {
    "submit": "提交", "retain": "保留", "on_request": "按要求提供",
    "not_applicable": "不适用", "": "尚未确认",
}


def _label(key: str) -> str:
    return _LABELS.get(str(key), str(key))


def _data_mode(task: Mapping[str, Any]) -> str:
    mode = str(task.get("data_mode") or "REAL").upper()
    if mode not in {"REAL", "TEST"}:
        raise ValueError("PPAP 数据模式必须为 REAL 或 TEST")
    return mode


def _text_lines(text: str, columns: int, *, font_size: int = 10, bold: bool = False) -> list[str]:
    """Conservative width-aware wrapping; every original character is retained.

    Excel widths use default-font digit units. CJK/full-width glyphs, wide Latin
    glyphs, word boundaries and explicit line breaks need separate accounting.
    This estimate reserves padding and font slack; it is not native auto-fit.
    """
    capacity = (columns * (_COLUMN_WIDTH * 7 + 5) - 12) / (font_size * 4 / 3)

    def width(char: str) -> float:
        if unicodedata.combining(char):
            return 0
        if char == "\t":
            return 2.8
        if unicodedata.east_asian_width(char) in {"W", "F", "A"}:
            result = 1.12
        elif char in "MWmw@%":
            result = 1.05
        elif char.isspace():
            result = 0.45
        else:
            result = 0.8 if char.isupper() else 0.7
        return result * (1.1 if bold else 1)

    lines: list[str] = []
    current = ""
    used = 0.0
    for token in re.findall(r"\r\n|\r|\n|[A-Za-z0-9_./:@%+\-]+|[^\r\n]", text):
        if token in {"\r\n", "\r", "\n"}:
            lines.append(current + token)
            current, used = "", 0.0
            continue
        token_width = sum(width(char) for char in token)
        # Normal Latin words move together; overlong identifiers wrap by glyph.
        if current and token_width <= capacity and used + token_width > capacity:
            lines.append(current)
            current, used = "", 0.0
        for char in token:
            char_width = width(char)
            if current and used + char_width > capacity:
                lines.append(current)
                current, used = "", 0.0
            current += char
            used += char_width
    lines.append(current)
    return lines


def _text_chunks(value: Any, columns: int, *, font_size: int = 10, bold: bool = False) -> list[Any]:
    if not isinstance(value, str):
        return [value]
    lines = _text_lines(value, columns, font_size=font_size, bold=bold)
    # A retained trailing newline adds an empty visual line at a chunk boundary.
    max_lines = max(1, int((_MAX_ROW_HEIGHT - 8) // (font_size * 1.8)) - 1)
    return ["".join(lines[i:i + max_lines]) for i in range(0, len(lines), max_lines)]


def _height(value: Any, columns: int, *, font_size: int = 10, bold: bool = False, minimum: int = 25) -> float:
    text = "" if value is None else str(value)
    lines = len(_text_lines(text, columns, font_size=font_size, bold=bold))
    height = max(minimum, math.ceil(8 + lines * font_size * 1.8))
    if height > _MAX_ROW_HEIGHT:
        raise ValueError("PPAP 文本必须先分段，不能超过安全行高")
    return height


def _assert_no_secrets(value: Any) -> None:
    if isinstance(value, Mapping):
        for key, child in value.items():
            if str(key).strip().lower() in _SECRET_KEYS:
                raise ValueError("PPAP 导出输入包含凭据字段，拒绝生成文件")
            _assert_no_secrets(child)
    elif isinstance(value, (list, tuple)):
        for child in value:
            _assert_no_secrets(child)


def _write_cell(ws: Any, row: int, col: int, value: Any, *, formula: bool = False) -> Any:
    cell = ws.cell(row, col)
    if formula:
        if not isinstance(value, str) or not value.startswith("="):
            raise ValueError("内部公式必须以 = 开头")
        cell.value = value
    elif value is None:
        cell.value = None
    elif isinstance(value, bool):
        cell.value = value
    elif isinstance(value, (int, float)):
        if not math.isfinite(float(value)):
            raise ValueError("PPAP 导出输入含非有限数值")
        cell.value = value
        cell.number_format = "General"
    else:
        text = str(value)
        if len(text) > 32767:
            raise ValueError("PPAP 单元格文本超过 Excel 限制，不做静默截断")
        cell.value = text
        # All caller-provided strings remain literal, including =HYPERLINK(...).
        cell.data_type = "s"
        cell.number_format = "@"
        if text.lstrip().startswith(("=", "+", "-", "@")):
            cell.quotePrefix = True
    cell.font = Font(name=_FONT, size=10, color=_NAVY)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = Border(bottom=Side(style="hair", color=_LINE))
    return cell


class _SheetWriter:
    """Render structured facts as native cells, never opaque JSON blobs."""

    def __init__(self, ws: Any, title: str, task: Mapping[str, Any], version: int):
        self.ws = ws
        self.row = 1
        self.data_mode = _data_mode(task)
        for col in range(1, MAX_COLUMNS + 1):
            self.ws.column_dimensions[get_column_letter(col)].width = _COLUMN_WIDTH
        if self.data_mode == "TEST":
            title = f"测试资料包：{title}"
        self._section(title, main_title=True)
        self._note(ppap_template.SHEET_BANNER, fill=_ORANGE)
        self._note(f"任务 {task.get('task_id', '')} · 零件 {task.get('part_no', '')} · "
                   f"侧别 {task.get('part_side', '')} · 图纸版本 {task.get('drawing_version', '')} · 包版本 {version}")
        self._note(_TEST_BANNER if self.data_mode == "TEST" else
                   "编制快照；生成不代表内部批准、客户接受、实物已提供或允许发运。",
                   fill=_ORANGE if self.data_mode == "TEST" else None)
        self.ws.freeze_panes = "A5"
        self.ws.sheet_view.showGridLines = False

    def _section(self, title: str, *, main_title: bool = False) -> None:
        size = 16 if main_title else 11
        for chunk in _text_chunks(title, MAX_COLUMNS, font_size=size, bold=True):
            for col in range(1, MAX_COLUMNS + 1):
                cell = _write_cell(self.ws, self.row, col, chunk if col == 1 else None)
                cell.fill = PatternFill("solid", fgColor=_NAVY if main_title else _TEAL)
                cell.font = Font(name=_FONT, size=size, color="FFFFFF", bold=True)
            self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=MAX_COLUMNS)
            self.ws.row_dimensions[self.row].height = _height(chunk, MAX_COLUMNS, font_size=size, bold=True,
                                                            minimum=34 if main_title else 26)
            self.row += 1

    def _note(self, text: Any, *, fill: str | None = None) -> None:
        text = str(text)
        # Split long explanatory fields into complete, printable pieces.
        chunks = _text_chunks(text, MAX_COLUMNS)
        for chunk in chunks:
            cell = _write_cell(self.ws, self.row, 1, chunk)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=MAX_COLUMNS)
            self.ws.row_dimensions[self.row].height = _height(chunk, MAX_COLUMNS, minimum=24)
            self.row += 1

    def fields(self, values: Mapping[str, Any]) -> None:
        for key, value in values.items():
            if isinstance(value, (Mapping, list, tuple)):
                self.render(_label(str(key)), value)
                continue
            labels = _text_chunks(_label(str(key)), 2)
            chunks = _text_chunks(value, MAX_COLUMNS - 2)
            for i in range(max(len(labels), len(chunks))):
                label = labels[i] if i < len(labels) else "（续）"
                chunk = chunks[i] if i < len(chunks) else None
                label_cell = _write_cell(self.ws, self.row, 1, label)
                label_cell.fill = PatternFill("solid", fgColor=_PALE)
                self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=2)
                _write_cell(self.ws, self.row, 3, chunk)
                self.ws.merge_cells(start_row=self.row, start_column=3, end_row=self.row, end_column=MAX_COLUMNS)
                self.ws.row_dimensions[self.row].height = max(_height(label, 2), _height(chunk, MAX_COLUMNS - 2))
                self.row += 1

    def table(self, rows: Sequence[Mapping[str, Any]], *, columns: list[str] | None = None) -> None:
        if not rows:
            self._note("未取得记录；留待补充，不按已完成或合格处理。", fill=_ORANGE)
            return
        keys = columns or list(dict.fromkeys(str(key) for row in rows for key, value in row.items()
                                            if not isinstance(value, (Mapping, list, tuple))))
        if not keys:
            for i, record in enumerate(rows, 1):
                self.render(f"记录 {i}", record)
            return
        identity_keys = [k for k in ("part_side", "op_id", "characteristic_id", "study_id", "point", "evidence_id") if k in keys]
        identity_keys = identity_keys[:2]
        other_keys = [k for k in keys if k not in identity_keys]
        width = MAX_COLUMNS - len(identity_keys)
        groups = [identity_keys + other_keys[i:i + width] for i in range(0, len(other_keys), width)] or [identity_keys]
        for g_index, group in enumerate(groups):
            if len(groups) > 1:
                self._note(f"字段分组 {g_index + 1}/{len(groups)}；对应相同记录顺序")
            self._table_header([_label(k) for k in group])
            for record in rows:
                values = {key: ("见本表下方展开内容" if isinstance(record.get(key), (Mapping, list, tuple))
                                else record.get(key)) for key in group}
                parts = {key: _text_chunks(value, 1) for key, value in values.items()}
                for continuation in range(max(len(v) for v in parts.values())):
                    max_height = 25
                    for col, key in enumerate(group, 1):
                        chunks = parts[key]
                        value = chunks[continuation] if continuation < len(chunks) else (
                            chunks[0] if key in identity_keys and len(chunks) == 1 else (
                                "（同一记录，见首行）" if key in identity_keys else None))
                        _write_cell(self.ws, self.row, col, value)
                        max_height = max(max_height, _height(value, 1))
                    self.ws.row_dimensions[self.row].height = max_height
                    self.row += 1
        for i, record in enumerate(rows, 1):
            nested = {str(k): v for k, v in record.items() if isinstance(v, (Mapping, list, tuple))}
            if nested:
                record_id = record.get("study_id") or record.get("op_id") or record.get("characteristic_id") or record.get("evidence_id") or i
                self.render(f"记录 {record_id} 展开", nested)

    def _table_header(self, columns: list[str]) -> None:
        parts = [_text_chunks(label, 1, bold=True) for label in columns]
        for continuation in range(max(map(len, parts), default=1)):
            height = 34
            for col, chunks in enumerate(parts, 1):
                label = chunks[continuation] if continuation < len(chunks) else None
                cell = _write_cell(self.ws, self.row, col, label)
                cell.font = Font(name=_FONT, size=10, bold=True, color=_NAVY)
                cell.fill = PatternFill("solid", fgColor=_PALE)
                height = max(height, _height(label, 1, bold=True, minimum=34))
            self.ws.row_dimensions[self.row].height = height
            self.row += 1

    def raw_values(self, values: list[Any], label: str) -> None:
        self._section(label)
        self._table_header(["样本顺序", "原始值"])
        start = self.row
        for i, value in enumerate(values, 1):
            for chunk in _text_chunks(value, 1):
                _write_cell(self.ws, self.row, 1, i)
                _write_cell(self.ws, self.row, 2, chunk)
                self.ws.row_dimensions[self.row].height = _height(chunk, 1, minimum=23)
                self.row += 1
        if values and all(isinstance(v, (int, float)) and not isinstance(v, bool) and math.isfinite(float(v)) for v in values):
            end = self.row - 1
            self._note("以下为可编辑 Excel 辅助统计；不替代系统确定性指标、规则确认或审核。")
            for title, fn in (("平均值", "AVERAGE"), ("最小值", "MIN"), ("最大值", "MAX")):
                _write_cell(self.ws, self.row, 1, title)
                self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=2)
                _write_cell(self.ws, self.row, 3, f"={fn}(B{start}:B{end})", formula=True)
                self.ws.row_dimensions[self.row].height = _height(title, 2, minimum=26)
                self.row += 1
            if len(values) >= 2:
                _write_cell(self.ws, self.row, 1, "样本标准差（n-1）")
                self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=2)
                _write_cell(self.ws, self.row, 3, f"=STDEV(B{start}:B{end})", formula=True)
                self.ws.row_dimensions[self.row].height = _height("样本标准差（n-1）", 2, minimum=26)
                self.row += 1

    def msa_measurements(self, measurements: list[Any]) -> None:
        self._section("MSA 原始测量网格")
        # Long-form layout supports any supplied trial count without hidden columns.
        self._table_header(["评价人", "零件标识", "重复测次", "原始实测值"])
        for measurement in measurements:
            if not isinstance(measurement, Mapping):
                self._note("记录格式无效；详见检查结果，不生成代替值。", fill=_ORANGE)
                continue
            for part in measurement.get("parts") or []:
                if not isinstance(part, Mapping):
                    self._note("零件测量记录格式无效。", fill=_ORANGE)
                    continue
                for trial, value in enumerate(part.get("trials") or [], 1):
                    items = (measurement.get("appraiser"), part.get("part") or part.get("part_id"), trial, value)
                    parts = [_text_chunks(item, 1) for item in items]
                    for continuation in range(max(map(len, parts))):
                        height = 24
                        for col, chunks in enumerate(parts, 1):
                            item = chunks[continuation] if continuation < len(chunks) else None
                            _write_cell(self.ws, self.row, col, item)
                            height = max(height, _height(item, 1, minimum=24))
                        self.ws.row_dimensions[self.row].height = height
                        self.row += 1

    def render(self, label: str, value: Any) -> None:
        if label == _label("values") and isinstance(value, list) and all(not isinstance(v, (Mapping, list, tuple)) for v in value):
            self.raw_values(value, label)
            return
        if label == _label("measurements") and isinstance(value, list):
            self.msa_measurements(value)
            return
        self._section(label)
        if isinstance(value, Mapping):
            if value:
                self.fields(value)
            else:
                self._note("尚无内容，需取得输入后编制。", fill=_ORANGE)
        elif isinstance(value, (list, tuple)):
            if not value:
                if label == _label("findings"):
                    self._note("本条目本次检查无发现；请结合整包检查与证据覆盖，不代表已批准。")
                elif label == _label("pending_confirmations"):
                    self._note("本次没有额外待确认事项，不代表已批准。")
                else:
                    self._note("尚无记录，需补充或确认适用性。", fill=_ORANGE)
            elif all(isinstance(v, Mapping) for v in value):
                self.table(value)
            else:
                for i, item in enumerate(value, 1):
                    if isinstance(item, (Mapping, list, tuple)):
                        self.render(f"记录 {i}", item)
                    else:
                        self.fields({str(i): item})
        else:
            self._note(value if value is not None else "未提供")

    def finish(self) -> None:
        ws = self.ws
        if self.row > 1_048_576:
            raise ValueError("PPAP 内容超出 Excel 单表行数限制")
        for col in range(1, MAX_COLUMNS + 1):
            ws.column_dimensions[get_column_letter(col)].width = _COLUMN_WIDTH
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_options.horizontalCentered = True
        ws.print_area = f"A1:H{self.row - 1}"
        ws.print_title_rows = "1:4"
        ws.page_margins.left = ws.page_margins.right = 0.25
        ws.page_margins.top = ws.page_margins.bottom = 0.4
        ws.oddFooter.center.text = "第 &P 页 / 共 &N 页"
        ws.oddFooter.right.text = "测试资料包（TEST）" if self.data_mode == "TEST" else "内部编制快照"


def _element_sheet(wb: Workbook, task: Mapping[str, Any], element: Mapping[str, Any], version: int) -> None:
    spec = ppap_template.element_spec(str(element.get("element_id")))
    if spec is None:
        raise ValueError("导出收到未声明的 PPAP 条目")
    writer = _SheetWriter(wb.create_sheet(spec.sheet_name), spec.title, task, version)
    writer.fields({"element_id": spec.element_id, "编制方式": spec.compilation_method,
                   "适用要求": _REQUIREMENTS.get(str(element.get("requirement") or ""), element.get("requirement")),
                   "适用依据": element.get("applicability_basis"), "条目状态": element.get("status") or "WAITING_DATA",
                   "状态说明": element.get("status_reason") or "尚无编制内容"})
    if not spec.generates_content:
        writer._note("本条目只归集原始文件；此表不是图纸、批准证明或实验室资质原件。", fill=_ORANGE)
    if spec.category == ppap_template.CAT_PHYSICAL:
        writer._note("此表仅登记实物及其证据；生成表格不代表样品、标准样件或检具已提供。", fill=_ORANGE)
    content = element.get("content") or element.get("content_json") or {}
    writer.render("编制内容", content)
    refs = element.get("source_refs") or element.get("source_refs_json") or {}
    writer.render("来源关联", refs)
    checks = element.get("check_result") or element.get("check_result_json") or {}
    writer.render("条目检查结果", checks)
    writer.finish()


def validate_workbook(data: bytes) -> dict[str, Any]:
    """Reopen the actual output and verify editable cells and printable structure.

    This confirms saved cell/formula/print definitions, not Excel's rendering engine
    or cached formula recalculation. No customer layout, images or OLE are claimed.
    """
    wb = load_workbook(BytesIO(data), data_only=False, keep_links=False)
    formulas: list[dict[str, Any]] = []
    auxiliary_checks: list[dict[str, Any]] = []
    sheets: list[dict[str, Any]] = []
    try:
        required = {"资料包概览", "条目目录", "缺项与冲突", "来源索引"} | {e.sheet_name for e in ppap_template.ELEMENTS}
        if set(wb.sheetnames) != required:
            raise ValueError("PPAP 工作簿缺少声明的内容表或出现额外工作表")
        for ws in wb:
            if any((dimension.height or 0) > _MAX_ROW_HEIGHT for dimension in ws.row_dimensions.values()):
                raise ValueError(f"工作表 {ws.title} 超出安全行高")
            if not ws.print_area or ws.print_title_rows != "$1:$4":
                raise ValueError(f"工作表 {ws.title} 缺少打印范围/重复表头")
            if ws.page_setup.fitToWidth != 1 or ws.page_setup.fitToHeight != 0:
                raise ValueError(f"工作表 {ws.title} 打印缩放错误")
            if ws.protection.sheet:
                raise ValueError(f"工作表 {ws.title} 被锁定，不满足可编辑要求")
            if ws["A2"].value != ppap_template.SHEET_BANNER:
                raise ValueError(f"工作表 {ws.title} 未标识项目自定义模板")
            populated = 0
            for row in ws:
                for cell in row:
                    if cell.value is not None:
                        populated += 1
                    if cell.data_type == "f":
                        formulas.append({"sheet": ws.title, "cell": cell.coordinate, "formula": cell.value})
                        match = re.fullmatch(r"=(AVERAGE|MIN|MAX|STDEV)\(B(\d+):B(\d+)\)", cell.value)
                        if not match:
                            raise ValueError("PPAP 辅助统计包含未声明公式")
                        fn, start, end = match.group(1), int(match.group(2)), int(match.group(3))
                        samples = [ws.cell(r, 2).value for r in range(start, end + 1)]
                        if not samples or not all(isinstance(v, (int, float)) and not isinstance(v, bool)
                                                  and math.isfinite(v) for v in samples):
                            raise ValueError("PPAP 辅助公式未对应完整有限数值样本")
                        if fn == "STDEV" and len(samples) < 2:
                            raise ValueError("PPAP 样本标准差缺少两个原始样本")
                        expected = {"AVERAGE": statistics.fmean, "MIN": min, "MAX": max,
                                    "STDEV": statistics.stdev}[fn](samples)
                        auxiliary_checks.append({"sheet": ws.title, "cell": cell.coordinate,
                                                 "source_range": f"B{start}:B{end}", "sample_count": len(samples),
                                                 "independent_expected_value": expected,
                                                 "method": "Python independent arithmetic; not native formula cache"})
                    if cell.data_type == "e":
                        raise ValueError(f"工作表 {ws.title} 含 Excel 错误单元格 {cell.coordinate}")
            sheets.append({"sheet": ws.title, "rows": ws.max_row, "columns": ws.max_column,
                           "populated_cells": populated, "print_area": str(ws.print_area),
                           "paper_size": ws.page_setup.paperSize, "orientation": ws.page_setup.orientation})
        return {"reopened": True, "sheet_count": len(sheets), "sheets": sheets,
                "formula_count": len(formulas), "formulas": formulas,
                "editable": True, "external_links": False,
                "formula_cache_recalculated": False,
                "auxiliary_formula_source_verified": True,
                "auxiliary_formula_checks": auxiliary_checks,
                "native_engine_verified": False,
                "max_row_height_points": _MAX_ROW_HEIGHT,
                "customer_layout_verified": False,
                "notice": "已回读内容、公式引用与打印定义，并独立计算辅助统计；公式缓存未重算，原生Excel分页渲染和正式客户模板尚未验证。"}
    finally:
        wb.close()


def export_package(*, task: dict[str, Any], elements: list[dict[str, Any]],
                   evidence: list[dict[str, Any]], issues: list[dict[str, Any]], version: int,
                   snapshot_hash: str, compile_run_id: str) -> dict[str, Any]:
    """Generate and reopen one package, saving only under controlled PPAP storage.

    Existing artifact paths are never overwritten, including approved versions.
    Idempotency is owned by the service transaction before invoking this function.
    """
    _assert_no_secrets({"task": task, "elements": elements, "evidence": evidence, "issues": issues})
    task_id = ppap_storage.safe_component(str(task.get("task_id") or ""), field="task_id")
    run_id = ppap_storage.safe_component(compile_run_id, field="compile_run_id")
    if not isinstance(version, int) or isinstance(version, bool) or version < 1:
        raise ValueError("产物版本必须为正整数")
    if not snapshot_hash or not isinstance(snapshot_hash, str):
        raise ValueError("导出必须指定输入快照指纹")
    element_map = {str(e.get("element_id")): e for e in elements}
    if len(element_map) != len(elements) or set(element_map) - set(ppap_template.ELEMENT_IDS):
        raise ValueError("PPAP 条目重复或包含未声明条目")
    template_source = task.get("template_source") or ppap_template.TEMPLATE_SOURCE
    data_mode = _data_mode(task)
    if template_source != ppap_template.TEMPLATE_SOURCE:
        raise ValueError("尚未实现正式客户模板映射，拒绝将内部模板标为客户模板")
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.title = "PPAP 测试资料包（TEST）" if data_mode == "TEST" else "PPAP 编制资料包"
    wb.properties.subject = ppap_template.SHEET_BANNER
    wb.properties.creator = "汽车零部件工厂智能体平台"
    wb.properties.description = (_TEST_BANNER + " " if data_mode == "TEST" else "") + "项目自定义模板；基于授权输入快照编制，不代表内部批准或客户接受。"
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True, forceFullCalc=True)
    cover = _SheetWriter(wb.create_sheet("资料包概览"), "PPAP 编制资料包", task, version)
    cover.fields({"data_mode": data_mode})
    cover.fields({k: task.get(k) for k in ("task_id", "customer_id", "customer_name", "part_no", "part_name",
                  "part_side", "drawing_version", "engineering_change_level", "submission_reason", "submission_level",
                  "owner", "reviewer", "due_date", "package_state", "customer_state")})
    cover.fields({"模板来源": template_source, "模板编号": task.get("template_id") or ppap_template.TEMPLATE_ID,
                  "模板版本": task.get("template_version") or ppap_template.TEMPLATE_VERSION,
                  "包版本": version, "产物状态": "DRAFT", "输入快照SHA256": snapshot_hash,
                  "编制运行": compile_run_id, "导出器版本": EXPORT_VERSION})
    cover._note(ppap_template.TEMPLATE_NOTICE, fill=_ORANGE)
    cover._note("可编辑文件是生成时的快照。离线修改不会自动更新系统数据、检查结果或批准版本；修改后需重新录入/确认输入并重新编制。辅助统计公式随 Excel 数据变化，系统指标仍以编制快照为准。")
    cover._note("未生成签字、客户批准或资质原件。客户处理状态与内部审核状态分别登记。")
    if task.get("reviewer") and task.get("reviewer") == task.get("owner"):
        cover._note("人员分离提示：任务负责人和指定审核人为同一账号。实际编制/审核身份以系统运行及审核记录为准；当前文件是未批准草稿。", fill=_ORANGE)
    cover.finish()
    directory = _SheetWriter(wb.create_sheet("条目目录"), "PPAP 条目目录", task, version)
    directory.table([{"条目编号": s.element_id, "名称": s.title, "适用要求": _REQUIREMENTS.get(str(element_map.get(s.element_id, {}).get("requirement") or ""), "尚未确认"),
                      "状态": element_map.get(s.element_id, {}).get("status") or "WAITING_DATA",
                      "编制方式": s.compilation_method, "责任岗位": s.responsible_role,
                      "状态说明": element_map.get(s.element_id, {}).get("status_reason") or "尚无内容"}
                     for s in ppap_template.ELEMENTS])
    directory.finish()
    for spec in ppap_template.ELEMENTS:
        _element_sheet(wb, task, element_map.get(spec.element_id, {"element_id": spec.element_id}), version)
    problems = _SheetWriter(wb.create_sheet("缺项与冲突"), "缺项与冲突清单", task, version)
    problems._note("只列本次收到的检查/待办事实；清单为空不证明证据读取完整或资料已批准。")
    problems.table(issues)
    if not issues:
        problems._note("本次未提供检查发现；请核对编制运行、条目适用性和证据覆盖。")
    problems.finish()
    sources = _SheetWriter(wb.create_sheet("来源索引"), "证据与来源索引", task, version)
    sources._note("定位信息为文字记录，不自动访问外部链接或执行嵌入对象。OK、EMPTY、ERROR 保持分别记录；ERROR 不等同于无此项。")
    sources.table(evidence, columns=[k for k in ("evidence_id", "source_type", "source_id", "source_version", "locator", "content_hash", "read_status", "review_state", "truncated", "summary", "element_ids", "error") if any(k in e for e in evidence)])
    sources.finish()
    output = BytesIO()
    wb.save(output)
    wb.close()
    data = output.getvalue()
    inspection = validate_workbook(data)
    file_name = f"ppap-v{version:04d}-{run_id}.xlsx"
    rel_path = f"{task_id}/v{version:04d}/{file_name}"
    target = ppap_storage.resolve_within(ppap_storage.ensure_storage_root(), rel_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    # 'xb' protects frozen versions and catches duplicate paths, instead of replace().
    with target.open("xb") as stream:
        stream.write(data)
    return {"rel_path": rel_path, "content_hash": ppap_storage.sha256_bytes(data), "media_type": MEDIA_TYPE,
            "byte_size": len(data), "file_name": file_name,
            "metadata": {"exporter_version": EXPORT_VERSION, "template_source": template_source,
                         "data_mode": data_mode, "test_package": data_mode == "TEST",
                         "is_customer_template": False, "artifact_state": "DRAFT", "version": version,
                         "snapshot_hash": snapshot_hash, "compile_run_id": compile_run_id,
                         "element_count": len(ppap_template.ELEMENTS), "source_count": len(evidence),
                         "issue_count": len(issues), "workbook_inspection": inspection}}
